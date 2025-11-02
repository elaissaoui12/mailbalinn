#!/usr/bin/env python3
"""
Professional CV Generator - Creates a Word document CV
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading_custom(doc, text, level=1, color=None):
    """Add a custom styled heading"""
    heading = doc.add_heading(text, level=level)
    if color:
        for run in heading.runs:
            run.font.color.rgb = color
    return heading

def add_section(doc, title, content_items):
    """Add a section with title and content"""
    # Section title
    heading = doc.add_heading(title, level=2)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 102, 204)  # Blue color
    
    # Section content
    for item in content_items:
        if isinstance(item, dict):
            # For structured items (like work experience)
            p = doc.add_paragraph()
            
            # Title/Position in bold
            if 'title' in item:
                run = p.add_run(item['title'])
                run.bold = True
                run.font.size = Pt(11)
            
            # Company/Institution and dates
            if 'organization' in item or 'dates' in item:
                p.add_run('\n')
                if 'organization' in item:
                    run = p.add_run(item['organization'])
                    run.italic = True
                if 'dates' in item:
                    run = p.add_run(f" | {item['dates']}")
                    run.font.color.rgb = RGBColor(128, 128, 128)
            
            # Description
            if 'description' in item:
                p.add_run(f"\n{item['description']}")
            
            # Bullet points
            if 'bullets' in item:
                for bullet in item['bullets']:
                    bullet_p = doc.add_paragraph(bullet, style='List Bullet')
                    bullet_p.paragraph_format.left_indent = Inches(0.5)
        else:
            # Simple text item
            doc.add_paragraph(item, style='List Bullet')
    
    doc.add_paragraph()  # Add spacing

def create_professional_cv():
    """Create a professional CV document"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Header - Name and Contact Info
    name = doc.add_heading('JOHN DOE', level=1)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in name.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0, 51, 102)
    
    # Contact Information
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_text = contact.add_run(
        'Email: john.doe@email.com | Phone: +1 (555) 123-4567 | '
        'LinkedIn: linkedin.com/in/johndoe | Location: New York, NY'
    )
    contact_text.font.size = Pt(10)
    contact_text.font.color.rgb = RGBColor(64, 64, 64)
    
    doc.add_paragraph()  # Spacing
    
    # Professional Summary
    add_section(doc, 'PROFESSIONAL SUMMARY', [
        'Results-driven professional with 5+ years of experience in software development and project management. '
        'Proven track record of delivering high-quality solutions and leading cross-functional teams. '
        'Strong expertise in full-stack development, agile methodologies, and client relationship management.'
    ])
    
    # Work Experience
    add_section(doc, 'WORK EXPERIENCE', [
        {
            'title': 'Senior Software Engineer',
            'organization': 'Tech Company Inc.',
            'dates': 'January 2021 - Present',
            'bullets': [
                'Led development of microservices architecture serving 1M+ daily active users',
                'Mentored team of 5 junior developers and conducted code reviews',
                'Improved application performance by 40% through optimization initiatives',
                'Collaborated with product team to define technical requirements and roadmap'
            ]
        },
        {
            'title': 'Software Engineer',
            'organization': 'Digital Solutions LLC',
            'dates': 'June 2019 - December 2020',
            'bullets': [
                'Developed and maintained full-stack web applications using React and Node.js',
                'Implemented CI/CD pipelines reducing deployment time by 60%',
                'Participated in agile ceremonies and sprint planning',
                'Created technical documentation and API specifications'
            ]
        },
        {
            'title': 'Junior Developer',
            'organization': 'StartUp Ventures',
            'dates': 'January 2018 - May 2019',
            'bullets': [
                'Built responsive web interfaces using modern JavaScript frameworks',
                'Collaborated with designers to implement pixel-perfect UI components',
                'Fixed bugs and implemented feature requests based on user feedback',
                'Participated in daily standups and team retrospectives'
            ]
        }
    ])
    
    # Education
    add_section(doc, 'EDUCATION', [
        {
            'title': 'Bachelor of Science in Computer Science',
            'organization': 'University of Technology',
            'dates': '2014 - 2018',
            'description': 'GPA: 3.8/4.0 | Dean\'s List | Relevant Coursework: Data Structures, Algorithms, Database Systems, Web Development'
        }
    ])
    
    # Skills
    add_section(doc, 'TECHNICAL SKILLS', [
        'Programming Languages: JavaScript, Python, Java, TypeScript, SQL',
        'Frontend: React, Vue.js, HTML5, CSS3, Tailwind CSS, Redux',
        'Backend: Node.js, Express, Django, FastAPI, RESTful APIs',
        'Databases: PostgreSQL, MongoDB, MySQL, Redis',
        'Tools & Technologies: Git, Docker, Kubernetes, AWS, CI/CD, Agile/Scrum',
        'Other: Unit Testing, System Design, Microservices, API Design'
    ])
    
    # Certifications
    add_section(doc, 'CERTIFICATIONS', [
        'AWS Certified Solutions Architect - Associate (2023)',
        'Professional Scrum Master I (PSM I) - Scrum.org (2022)',
        'MongoDB Certified Developer (2021)'
    ])
    
    # Projects (Optional)
    add_section(doc, 'KEY PROJECTS', [
        {
            'title': 'E-Commerce Platform',
            'description': 'Built a scalable e-commerce platform handling 10K+ transactions daily',
            'bullets': [
                'Technologies: React, Node.js, PostgreSQL, Redis, AWS',
                'Implemented payment gateway integration and inventory management system',
                'Achieved 99.9% uptime with automated monitoring and alerting'
            ]
        },
        {
            'title': 'Real-Time Analytics Dashboard',
            'description': 'Developed real-time data visualization dashboard for business intelligence',
            'bullets': [
                'Technologies: Vue.js, Python, WebSockets, D3.js',
                'Processed and visualized 1M+ data points in real-time',
                'Reduced data processing latency by 70%'
            ]
        }
    ])
    
    # Save the document
    doc.save('/vercel/sandbox/Professional_CV.docx')
    print("✓ CV created successfully: Professional_CV.docx")
    print("\nYour professional CV has been generated!")
    print("You can now customize it with your own information.")

if __name__ == '__main__':
    try:
        create_professional_cv()
    except ImportError:
        print("Error: python-docx library is required.")
        print("Installing python-docx...")
        import subprocess
        subprocess.check_call(['pip', 'install', 'python-docx'])
        print("\nLibrary installed. Running CV generator...")
        create_professional_cv()
